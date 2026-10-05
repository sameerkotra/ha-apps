"""Sheets (SPEC §7.2, §8, §17.10): .xlsx and .csv round trips (formulas with cached values, formats, widths,
frozen panes, tabs, charts, conditional colours, the totals row), the read-only rule for .xlsx files made
elsewhere (macros, pivot tables, pictures, comments, merged cells, links to other workbooks, data validation,
charts, unknown functions …), malicious files (zip bombs, XML entities, huge sheets — refused safely and
quickly), saving changed cells with per-cell conflicts, limits, export, import, Convert to .xlsx, Save a copy to
edit, search of cells (results open at the cell), and the personal setting for new sheets."""
import _env  # noqa: F401

import io
import json
import os
import re
import subprocess
import sys
import time
import unittest
import zipfile

from base import DEV, KABIR, MEERA, ApiBase, read, write
from app.formats import sheet_csv, sheet_model as M, sheet_xlsx

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CUR = {"f": "currency", "d": 2}


def grid(name="Sheet1", cells=None, **kw):
    t = {"name": name, "kind": "grid", "cells": cells or {}, "cols": {}, "freeze": {"r": 0, "c": 0}, "totals": None,
         "cond": [], "charts": []}
    t.update(kw)
    return t


def budget():
    return {"tabs": [
        grid("Budget", {
            "A1": {"v": "Item", "b": 1}, "B1": {"v": "Cost", "b": 1, "al": "right"},
            "A2": {"v": "Rent"}, "B2": dict(CUR, v=1200),
            "A3": {"v": "Food"}, "B3": dict(CUR, v=-35.5, red=1),
            "A4": {"v": "Car"}, "B4": dict(CUR, v="='Car costs'!B4*2", c=180),
            "C2": {"v": '=CONCAT(A2," x")', "c": "Rent x"}, "C3": {"v": "=B2/0", "c": {"e": "#DIV/0!"}},
            "C4": {"v": "=B2>100", "c": True}, "D2": {"v": 46000, "f": "date"}, "D3": {"v": 0.25, "f": "percent", "d": 1},
            "E2": {"v": "=not a formula", "q": 1}, "E3": {"v": "=DAYS(D2,D2-7)", "c": 7},
            "E4": {"v": '=XLOOKUP("Car",A2:A4,B2:B4)', "c": 180}},
             cols={"A": 160, "B": 110}, freeze={"r": 1, "c": 1}, totals={"B": "SUM"},
             cond=[{"range": "B2:B4", "type": "gt", "a": 300, "fill": "red", "text": None},
                   {"range": "A2:A4", "type": "contains", "a": "ar", "fill": None, "text": "blue"},
                   {"range": "D2:D4", "type": "before", "a": 46100, "fill": "green", "text": None},
                   {"range": "B2:B4", "type": "top", "n": 1, "fill": "amber", "text": None},
                   {"range": "A2:A9", "type": "dup", "fill": "grey", "text": None},
                   {"range": "B2:B4", "type": "between", "a": 1, "b": 500, "fill": "purple", "text": "red"},
                   {"range": "B2:B4", "type": "lt", "a": 0, "fill": None, "text": "red"},
                   {"range": "D2:D4", "type": "after", "a": 45000, "fill": "blue", "text": None},
                   {"range": "B2:B4", "type": "bottom", "n": 2, "fill": "grey", "text": None}],
             charts=[{"id": "c1", "type": "bar", "range": "A1:B4", "title": "Costs"},
                     {"id": "c2", "type": "line", "range": "A1:B4", "title": ""}]),
        grid("Car costs", {"A4": {"v": "Hire"}, "B4": dict(CUR, v=90)}),
        {"name": "Share", "kind": "chart", "chart": {"type": "pie", "title": "Share", "source": {"tab": "Budget", "range": "A1:B4"}}},
    ]}


def zip_with(data: bytes, extra: dict, drop=()) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in zin.infolist():
            if info.filename in drop:
                continue
            raw = zin.read(info.filename)
            if info.filename in extra and callable(extra[info.filename]):
                raw = extra[info.filename](raw)
            z.writestr(info, raw)
        for name, raw in extra.items():
            if not callable(raw) and name not in zin.namelist():
                z.writestr(name, raw)
    return out.getvalue()


def foreign(make=None) -> bytes:
    """An .xlsx made by "another app" (openpyxl without the app's property)."""
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws["A1"], ws["B1"], ws["A2"], ws["B2"], ws["B3"] = "Bill", "Amount", "Power", 80, "=SUM(B2:B2)"
    if make:
        make(wb, ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class Formats(unittest.TestCase):
    """formats/ without the API."""

    def test_xlsx_round_trip_keeps_everything_the_app_has(self):
        sheet = M.clean_sheet(budget())
        data = sheet_xlsx.write(sheet, "EUR")
        out = sheet_xlsx.read(data)
        self.assertFalse(out["readOnly"], out["features"])
        self.assertTrue(out["owned"])
        back = out["sheet"]
        self.assertEqual([t["name"] for t in back["tabs"]], ["Budget", "Car costs", "Share"])
        b, want = back["tabs"][0], sheet["tabs"][0]
        self.assertEqual({k: M.content(v) for k, v in b["cells"].items()}, {k: M.content(v) for k, v in want["cells"].items()})
        for ref in ("B4", "C2", "C3", "C4", "E3", "E4"):                     # cached values came back
            self.assertEqual(b["cells"][ref]["c"], want["cells"][ref]["c"], ref)
        self.assertEqual(b["cols"]["A"], 160)
        self.assertEqual(b["freeze"], {"r": 1, "c": 1})
        self.assertEqual(b["totals"], {"B": "SUM"})                         # recognised, taken out of the grid
        self.assertNotIn("B5", b["cells"])
        self.assertEqual(b["charts"], want["charts"])
        key = lambda r: json.dumps(r, sort_keys=True)
        self.assertEqual(sorted(map(key, b["cond"])), sorted(map(key, want["cond"])))
        self.assertEqual(back["tabs"][2], sheet["tabs"][2])

    def test_xlsx_as_excel_sees_it(self):
        """Opened with openpyxl in a fresh process: formulas (with Excel's prefixes), cached values, formats."""
        data = sheet_xlsx.write(M.clean_sheet(budget()), "EUR")
        path = os.path.join(os.environ["DATA_DIR"], "probe.xlsx")
        with open(path, "wb") as f:
            f.write(data)
        code = f"""
import json
from openpyxl import load_workbook
wb = load_workbook({path!r}); wv = load_workbook({path!r}, data_only=True)
ws, vs = wb["Budget"], wv["Budget"]
print(json.dumps({{"f": {{c: ws[c].value for c in ("B4", "C2", "E3", "E4", "B5", "A5")}},
  "v": {{c: vs[c].value for c in ("B4", "C2", "C3", "C4", "E3", "E4", "B5")}},
  "fmt": [ws["B2"].number_format, ws["B3"].number_format, ws["D2"].number_format, ws["D3"].number_format],
  "bold": ws["A1"].font.b, "align": ws["B1"].alignment.horizontal, "width": ws.column_dimensions["A"].width,
  "freeze": ws.freeze_panes, "e2": [ws["E2"].value, ws["E2"].data_type, ws["E2"].quotePrefix],
  "charts": [len(ws._charts), len(wb["Share"]._charts)], "sheets": wb.sheetnames, "cf": sum(len(r.rules) for r in ws.conditional_formatting),
  "props": {{p.name: p.value for p in wb.custom_doc_props}}.get("HouseholdDocs"), "calc": wb.calculation.fullCalcOnLoad}}, default=str))
"""
        r = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, timeout=60)
        info = json.loads(r.stdout)
        self.assertEqual(info["f"]["C2"], '=_xlfn.CONCAT(A2," x")')
        self.assertEqual(info["f"]["E3"], "=_xlfn.DAYS(D2,D2-7)")
        self.assertEqual(info["f"]["E4"], '=_xlfn.XLOOKUP("Car",A2:A4,B2:B4)')
        self.assertEqual(info["f"]["B4"], "='Car costs'!B4*2")
        self.assertEqual(info["f"]["B5"], "=SUM(B1:B4)")
        self.assertEqual(info["f"]["A5"], "Total")
        self.assertEqual(info["v"], {"B4": 180, "C2": "Rent x", "C3": "#DIV/0!", "C4": True, "E3": 7, "E4": 180, "B5": 1344.5})
        self.assertEqual(info["fmt"], ['"€"#,##0.00', '"€"#,##0.00;[Red]-"€"#,##0.00', "yyyy-mm-dd", "0.0%"])
        self.assertIs(info["bold"], True)
        self.assertEqual(info["align"], "right")
        self.assertAlmostEqual(info["width"], 22.14, places=1)
        self.assertEqual(info["freeze"], "B2")
        self.assertEqual(info["e2"], ["=not a formula", "s", True])
        self.assertEqual(info["charts"], [2, 1])
        self.assertEqual(info["sheets"], ["Budget", "Car costs", "Share"])
        self.assertEqual(info["cf"], 9)
        self.assertEqual(info["props"], "1")
        self.assertIs(info["calc"], True)

    def test_number_formats_from_excel(self):
        f = M.format_from_excel
        self.assertEqual(f("General"), {})
        self.assertEqual(f("0.00"), {"f": "number", "d": 2})
        self.assertEqual(f("#,##0"), {"f": "number", "d": 0})
        self.assertEqual(f("0%"), {"f": "percent", "d": 0})
        self.assertEqual(f('[$€-x-euro2] #,##0.00'), {"f": "currency", "d": 2})
        self.assertEqual(f('"$"#,##0.00_);[Red]("$"#,##0.00)'), {"f": "currency", "d": 2, "red": 1})
        self.assertEqual(f("yyyy-mm-dd"), {"f": "date"})
        self.assertEqual(f("d/m/yy"), {"f": "date"})
        self.assertEqual(f("mmm-yy"), {"f": "date"})
        self.assertEqual(f("@"), {"f": "text"})
        self.assertIsNone(f("hh:mm"))
        self.assertIsNone(f("0.00E+00"))
        self.assertIsNone(f("# ?/?"))
        for cell in ({"f": "number", "d": 3}, {"f": "currency", "d": 0, "red": 1}, {"f": "percent", "d": 2}, {"f": "date"},
                     {"f": "text"}, {"red": 1}):
            self.assertEqual(f(M.excel_format(cell, "USD")), cell, cell)

    def test_csv_dialects_round_trip(self):
        cases = [
            b"Item,Cost\nRent,1200\nFood,35.5\n,=SUM(B2:B3)\n",
            b"Item;Cost\r\nRent;1200\r\nCaf\xc3\xa9;007\r\n",
            b"\xef\xbb\xbfName\tWhen\nAsha\t2026-03-09\nKabir\t1.50\n",
            b'Quote,Text\n"a,b","say ""hi"""\nTRUE,true\n',
        ]
        for data in cases:
            r = sheet_csv.read(data)
            again = sheet_csv.write(r["sheet"], r["dialect"])
            self.assertEqual(again, data, data)
        r = sheet_csv.read("Ä;Ö\n1;2\n".encode("utf-16"))
        self.assertEqual(r["dialect"]["sep"], ";")
        self.assertEqual(r["sheet"]["tabs"][0]["cells"]["A1"]["v"], "Ä")
        self.assertTrue(sheet_csv.write(r["sheet"], r["dialect"]).startswith(b"\xef\xbb\xbf\xc3\x84;"))   # → UTF-8 with a BOM
        cells = sheet_csv.read(b"a,b\n1,2.5\n0.10,-3\n2026-02-30,x\n")["sheet"]["tabs"][0]["cells"]
        self.assertEqual(cells["A2"], {"v": 1})
        self.assertEqual(cells["A3"], {"v": "0.10"})               # wouldn't come back the same as a number
        self.assertEqual(cells["B3"], {"v": -3})
        self.assertEqual(cells["A4"], {"v": "2026-02-30"})          # not a real date

    def test_csv_limits(self):
        with self.assertRaises(sheet_csv.CsvError):
            sheet_csv.read(b"x\n" * 5001 + b"y\n")
        with self.assertRaises(sheet_csv.CsvError):
            sheet_csv.read((",".join(["1"] * 101)).encode())

    def test_function_lists_match_the_engine(self):
        js = read(os.path.join(HERE, "app", "static", "sheetcalc.js"))
        names = set(re.findall(r'\bdef\("([A-Z]+)"', js))
        self.assertEqual(names, set(M.KNOWN_FUNCTIONS))

    def test_formula_helpers(self):
        self.assertEqual(M.functions_in('=SUM(A1)+IF(B1,"MAX(1)",_xlfn.XLOOKUP(1,A:A,B:B))'), {"SUM", "IF", "XLOOKUP"})
        self.assertEqual(M.to_excel_formula('CONCAT("CONCAT(",DAYS(A1,B1))'), '_xlfn.CONCAT("CONCAT(",_xlfn.DAYS(A1,B1))')
        self.assertEqual(M.from_excel_formula("_xlfn.CONCAT(A1)"), "CONCAT(A1)")
        self.assertTrue(M.can_calculate("SUM(A1:A3)*2"))
        self.assertFalse(M.can_calculate("SQRT(4)"))
        self.assertFalse(M.can_calculate("SUM(Table1[Cost])"))
        self.assertFalse(M.can_calculate("SUM({1,2})"))
        self.assertFalse(M.can_calculate("Rate*2", {"Rate"}))
        self.assertTrue(M.can_calculate('"Rate"&A1', {"Rate"}))


class ForeignFiles(unittest.TestCase):
    """The read-only rule (SPEC §8.4): anything the app can't keep → read only, with the list."""

    def features(self, data):
        return sheet_xlsx.read(data)

    def test_plain_foreign_file_opens_for_editing(self):
        out = self.features(foreign())
        self.assertFalse(out["readOnly"], out["features"])
        self.assertFalse(out["owned"])
        cells = out["sheet"]["tabs"][0]["cells"]
        self.assertEqual(cells["B3"]["v"], "=SUM(B2:B2)")

    def test_each_feature_makes_it_read_only(self):
        from openpyxl.chart import BarChart, Reference
        from openpyxl.comments import Comment
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.workbook.defined_name import DefinedName

        def chart(wb, ws):
            c = BarChart()
            c.add_data(Reference(ws, min_col=2, min_row=1, max_row=2), titles_from_data=True)
            ws.add_chart(c, "D2")

        def dv(wb, ws):
            v = DataValidation(type="list", formula1='"a,b"')
            ws.add_data_validation(v)
            v.add("C1")

        def names(wb, ws):
            wb.defined_names["Rate"] = DefinedName("Rate", attr_text="Bills!$B$2")
            ws["C1"] = "=Rate*2"

        def fill(wb, ws):
            from openpyxl.styles import PatternFill
            ws["A1"].fill = PatternFill("solid", start_color="FFFF00")

        def hidden(wb, ws):
            wb.create_sheet("Secret").sheet_state = "hidden"

        cases = {
            "Merged cells": lambda wb, ws: ws.merge_cells("A5:B5"),
            "Comments or notes": lambda wb, ws: setattr(ws["A1"], "comment", Comment("note", "someone")),
            "Data validation (drop-down lists, checks)": dv,
            "Charts not made by the app": chart,
            "Functions the app doesn't have: SQRT": lambda wb, ws: ws.__setitem__("C1", "=SQRT(4)"),
            "Named ranges": names,
            "Colours, borders or fonts the app doesn't keep": fill,
            "Hidden tabs": hidden,
            "Number formats the app doesn't have (times, scientific …)": lambda wb, ws: setattr(ws["B2"], "number_format", "hh:mm"),
            "Filters (AutoFilter)": lambda wb, ws: setattr(ws.auto_filter, "ref", "A1:B2"),
        }
        for label, make in cases.items():
            out = self.features(foreign(make))
            self.assertTrue(out["readOnly"], label)
            self.assertIn(label, out["features"], (label, out["features"]))

    def test_parts_that_make_it_read_only(self):
        base = foreign()
        png = b"\x89PNG\r\n\x1a\n" + b"\0" * 20
        cases = {
            "Macros (VBA)": {"xl/vbaProject.bin": b"\xd0\xcf\x11\xe0 fake macros"},
            "Pivot tables": {"xl/pivotTables/pivotTable1.xml": b"<pivotTableDefinition/>",
                             "xl/pivotCache/pivotCacheDefinition1.xml": b"<x/>"},
            "Pictures": {"xl/media/image1.png": png},
            "Links to other workbooks": {"xl/externalLinks/externalLink1.xml": b"<externalLink/>"},
            "Excel tables": {"xl/tables/table1.xml": b"<table/>"},
            "Slicers or timelines": {"xl/slicers/slicer1.xml": b"<x/>"},
            "Other Excel features": {"xl/somethingNew/part.xml": b"<x/>"},
        }
        for label, extra in cases.items():
            out = self.features(zip_with(base, extra))
            self.assertTrue(out["readOnly"], label)
            self.assertIn(label, out["features"], (label, out["features"]))
        # a macro-enabled content type alone
        ct = lambda raw: raw.replace(b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
                                     b"application/vnd.ms-excel.sheet.macroEnabled.main+xml")
        out = self.features(zip_with(base, {"[Content_Types].xml": ct}))
        self.assertIn("Macros (VBA)", out["features"])
        self.assertEqual(out["sheet"]["tabs"][0]["cells"]["A2"], {"v": "Power"})          # still shown

    def test_app_file_with_a_chart_added_elsewhere_is_read_only(self):
        """The app's own file, but someone added a chart in Excel: not the app's chart → read only."""
        from openpyxl import load_workbook
        from openpyxl.chart import LineChart, Reference
        data = sheet_xlsx.write(M.clean_sheet({"tabs": [grid("S", {"A1": {"v": 1}, "A2": {"v": 2}})]}))
        self.assertFalse(sheet_xlsx.read(data)["readOnly"])
        wb = load_workbook(io.BytesIO(data))
        c = LineChart()
        c.add_data(Reference(wb["S"], min_col=1, min_row=1, max_row=2))
        wb["S"].add_chart(c, "C3")
        buf = io.BytesIO()
        wb.save(buf)
        out = sheet_xlsx.read(buf.getvalue())
        self.assertTrue(out["readOnly"])
        self.assertIn("Charts not made by the app", out["features"])


class Malicious(unittest.TestCase):
    """Untrusted .xlsx files are refused safely and quickly (SPEC §3.4 risk 5)."""

    def refused(self, data, words):
        t = time.time()
        with self.assertRaises(sheet_xlsx.ReadError) as cm:
            sheet_xlsx.read(data)
        took = time.time() - t
        self.assertLess(took, 15, f"took {took:.1f} s")
        self.assertRegex(str(cm.exception), words)
        return took

    def test_zip_bomb(self):
        z = io.BytesIO()
        with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
            zf.writestr("[Content_Types].xml", "<Types/>")
            zf.writestr("xl/workbook.xml", "<workbook/>")
            zf.writestr("xl/worksheets/sheet1.xml", b"\0" * (80 * 1024 * 1024))
        self.assertLess(len(z.getvalue()), 1024 * 1024)
        self.refused(z.getvalue(), "zip bomb")

    def test_xml_entity_expansion(self):
        ents = '<!ENTITY lol "lol">' + "".join(
            f'<!ENTITY lol{i} "{("&lol%s;" % (i - 1 if i > 1 else "")) * 10}">' for i in range(1, 10))
        bomb = f'<?xml version="1.0"?><!DOCTYPE lolz [{ents}]>'.encode()

        def put(raw):
            body = raw.split(b"?>", 1)[1] if raw.startswith(b"<?xml") else raw
            return bomb + body.replace(b"<v>", b"<v>&lol9;", 1)
        for part in ("xl/worksheets/sheet1.xml", "xl/workbook.xml", "xl/styles.xml"):
            self.refused(zip_with(foreign(), {part: put}), "entities or a DTD")
        ext = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]>'
        self.refused(zip_with(foreign(), {"xl/worksheets/sheet1.xml": lambda raw: ext + raw.split(b"?>", 1)[-1]}), "entities or a DTD")
        # past the early check (a DTD after 64 KB of whitespace): defusedxml still refuses it
        late = b'<?xml version="1.0"?>' + b" " * 70000 + bomb.split(b"?>", 1)[1]
        self.refused(zip_with(foreign(), {"xl/worksheets/sheet1.xml": lambda raw: late + raw.split(b"?>", 1)[-1]}),
                     "entities or a DTD|couldn't be read")

    def test_huge_sheets(self):
        cols = [M.col_name(i) for i in range(100)]
        rows = "".join(f'<row r="{r}">' + "".join(f'<c r="{c}{r}"><v>1</v></c>' for c in cols) + "</row>"
                       for r in range(1, 2_002))
        big = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + rows
               + "</sheetData></worksheet>")
        took = self.refused(zip_with(foreign(), {"xl/worksheets/sheet1.xml": lambda raw: big.encode()}), "200,000 filled cells")
        far = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
               '<row r="9000"><c r="A9000"><v>1</v></c></row></sheetData></worksheet>')
        self.refused(zip_with(foreign(), {"xl/worksheets/sheet1.xml": lambda raw: far.encode()}), "5,000 rows")
        wide = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
                '<row r="1"><c r="ZZ1"><v>1</v></c></row></sheetData></worksheet>')
        self.refused(zip_with(foreign(), {"xl/worksheets/sheet1.xml": lambda raw: wide.encode()}), "100 columns")
        self.assertLess(took, 15)

    def test_too_many_tabs_garbage_and_encrypted(self):
        from openpyxl import Workbook

        def many(wb, ws):
            for i in range(10):
                wb.create_sheet(f"T{i}")
        self.refused(foreign(many), "at most 10")
        self.refused(b"PK\x03\x04 not really a zip", "isn't an .xlsx")
        with self.assertRaises(sheet_xlsx.ReadError):
            sheet_xlsx.read(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 600)     # an encrypted (OLE) workbook
        _ = Workbook

    def test_the_worker_is_limited_and_killed(self):
        """A worker that runs too long is killed at the time limit."""
        data = foreign()
        t = time.time()
        with self.assertRaises(sheet_xlsx.ReadError) as cm:
            sheet_xlsx.run_worker(data, timeout=0.01)
        self.assertLess(time.time() - t, 5)
        self.assertIn("stopped", str(cm.exception))


class SheetApi(ApiBase):
    def make(self, name="Budget", sheet=None, h=KABIR, fmt=None, parent=None):
        body = {"kind": "sheet", "name": name, "parentId": parent}
        if fmt:
            body["format"] = fmt
        it = self.ok(self.post("/api/docs", body, h), 201)
        if sheet is not None:
            etag = self.open(it["id"], h)["etag"]
            self.ok(self.patch(f"/api/docs/{it['id']}", {"etag": etag, "sheet": sheet}, h))
        return it

    def cells(self, nid, h=KABIR, tab=0):
        return self.open(nid, h)["sheet"]["tabs"][tab]["cells"]

    def change(self, nid, changes, etag, h=KABIR, tabs=None):
        body = {"etag": etag, "cells": [{"tab": t, "ref": r, "was": w, "now": n} for t, r, w, n in changes]}
        if tabs is not None:
            body["tabs"] = tabs
        return self.patch(f"/api/docs/{nid}", body, h)

    # -- create, open, save ------------------------------------------------------
    def test_new_sheet_is_an_xlsx_the_app_owns(self):
        it = self.make()
        self.assertEqual(it["name"], "Budget.xlsx")
        self.assertEqual(it["kind"], "sheet")
        doc = self.open(it["id"])
        self.assertEqual(doc["format"], "xlsx")
        self.assertFalse(doc["readOnly"])
        self.assertTrue(doc["canEdit"])
        self.assertEqual(doc["sheet"], M.empty_sheet())
        self.assertEqual(doc["limits"], {"tabs": 10, "rows": 5000, "cols": 100, "cells": 200000})
        with open(self.real(it["id"]), "rb") as f:
            self.assertTrue(f.read(2) == b"PK")
        c = self.make("Plain", fmt="csv")
        self.assertEqual(c["name"], "Plain.csv")
        self.assertEqual(self.open(c["id"])["format"], "csv")
        self.assertEqual(self.post("/api/docs", {"kind": "note", "name": "x", "format": "csv"}).status_code, 422)

    def test_whole_sheet_save_and_reopen(self):
        it = self.make("Budget", budget())
        doc = self.open(it["id"])
        self.assertEqual([t["name"] for t in doc["sheet"]["tabs"]], ["Budget", "Car costs", "Share"])
        b = doc["sheet"]["tabs"][0]
        self.assertEqual(b["totals"], {"B": "SUM"})
        self.assertEqual(len(b["cond"]), 9)
        self.assertEqual(b["cells"]["B4"]["c"], 180)

    def test_changed_cells_merge_and_only_the_same_cell_conflicts(self):
        it = self.make("Shared", {"tabs": [grid("S", {"A1": {"v": 1}, "B1": {"v": 2}})]})
        self.ok(self.share(it["id"], MEERA, "editor"))
        mine = self.open(it["id"])
        theirs = self.open(it["id"], MEERA)
        self.ok(self.change(it["id"], [("S", "A1", {"v": 1}, {"v": 10})], theirs["etag"], MEERA))
        # a different cell, from the older etag: saved, and the answer carries Meera's change
        out = self.ok(self.change(it["id"], [("S", "B1", {"v": 2}, {"v": 20})], mine["etag"]))
        self.assertEqual(out["sheet"]["tabs"][0]["cells"], {"A1": {"v": 10}, "B1": {"v": 20}})
        self.assertEqual(out["by"], "Meera Rao")
        self.assertEqual(self.cells(it["id"]), {"A1": {"v": 10}, "B1": {"v": 20}})
        # the same cell: 409 for that cell; the other change in the request is still saved
        r = self.change(it["id"], [("S", "A1", {"v": 1}, {"v": 99}), ("S", "C1", None, {"v": "new"})], mine["etag"])
        self.assertEqual(r.status_code, 409, r.text)
        d = r.json()["detail"]
        self.assertEqual([(c["tab"], c["ref"], c["theirs"]) for c in d["conflicts"]], [("S", "A1", {"v": 10})])
        self.assertEqual(self.cells(it["id"]), {"A1": {"v": 10}, "B1": {"v": 20}, "C1": {"v": "new"}})
        # Keep mine: send again with what's there now as "was"
        self.ok(self.change(it["id"], [("S", "A1", {"v": 10}, {"v": 99})], d["etag"]))
        self.assertEqual(self.cells(it["id"])["A1"], {"v": 99})
        # the same change made twice isn't a conflict
        self.ok(self.change(it["id"], [("S", "B1", {"v": 2}, {"v": 20})], mine["etag"]))
        # a cell in a tab that's gone
        r = self.change(it["id"], [("Gone", "A1", None, {"v": 1})], mine["etag"])
        self.assertEqual(r.status_code, 409)
        self.assertTrue(r.json()["detail"]["conflicts"][0]["gone"])

    def test_format_changes_conflict_too_and_computed_values_dont(self):
        it = self.make("F", {"tabs": [grid("S", {"A1": {"v": 5}, "A2": {"v": "=A1*2", "c": 10}})]})
        self.ok(self.share(it["id"], MEERA, "editor"))
        old = self.open(it["id"])["etag"]
        self.ok(self.change(it["id"], [("S", "A1", {"v": 5}, {"v": 5, "b": 1})], old, MEERA))
        r = self.change(it["id"], [("S", "A1", {"v": 5}, {"v": 5, "f": "currency"})], old)
        self.assertEqual(r.status_code, 409)
        # a formula's new value (it reads a changed cell): no conflict, the value is stored
        now = self.open(it["id"])["etag"]
        self.ok(self.change(it["id"], [("S", "A2", {"v": "=A1*2", "c": 10}, {"v": "=A1*2", "c": 12})], now))
        self.assertEqual(self.cells(it["id"])["A2"], {"v": "=A1*2", "c": 12})

    def test_shape_changes_need_an_unchanged_file(self):
        it = self.make("Shape", {"tabs": [grid("S", {"A1": {"v": 1}})]})
        self.ok(self.share(it["id"], MEERA, "editor"))
        old = self.open(it["id"])["etag"]
        self.ok(self.change(it["id"], [("S", "A2", None, {"v": 2})], old, MEERA))
        r = self.patch(f"/api/docs/{it['id']}", {"etag": old, "sheet": {"tabs": [grid("S", {"A2": {"v": 1}})]}})
        self.assertEqual(r.status_code, 409)
        d = r.json()["detail"]
        self.assertTrue(d["stale"])
        self.assertEqual(d["sheet"]["tabs"][0]["cells"], {"A1": {"v": 1}, "A2": {"v": 2}})
        self.ok(self.patch(f"/api/docs/{it['id']}", {"etag": d["etag"], "sheet": {"tabs": [grid("S", {"A3": {"v": 1}})]}}))

    def test_tab_settings_save_without_conflicts(self):
        it = self.make("Meta", {"tabs": [grid("S", {"A1": {"v": 1}})]})
        etag = self.open(it["id"])["etag"]
        meta = {"name": "S", "cols": {"A": 200}, "freeze": {"r": 1, "c": 0}, "totals": {"A": "MAX"},
                "cond": [{"range": "A1:A9", "type": "dup", "fill": "red", "text": None}],
                "charts": [{"id": "x", "type": "pie", "range": "A1:A3", "title": "T"}]}
        self.ok(self.change(it["id"], [], etag, tabs=[meta]))
        t = self.open(it["id"])["sheet"]["tabs"][0]
        self.assertEqual((t["cols"], t["freeze"], t["totals"], len(t["cond"]), t["charts"][0]["id"]),
                         ({"A": 200}, {"r": 1, "c": 0}, {"A": "MAX"}, 1, "x"))

    def test_outside_change_then_save_merges_and_keeps_a_version(self):
        it = self.make("Out", {"tabs": [grid("S", {"A1": {"v": 1}})]})
        etag = self.open(it["id"])["etag"]
        data = sheet_xlsx.write(M.clean_sheet({"tabs": [grid("S", {"A1": {"v": 1}, "Z9": {"v": "outside"}})]}))
        time.sleep(0.01)
        write(self.real(it["id"]), data)
        out = self.ok(self.change(it["id"], [("S", "B1", None, {"v": 2})], etag))
        self.assertTrue(out["outside"])
        self.assertEqual(self.cells(it["id"]), {"A1": {"v": 1}, "B1": {"v": 2}, "Z9": {"v": "outside"}})
        vs = self.ok(self.get(f"/api/docs/{it['id']}/versions"))["versions"]
        self.assertTrue(any(v["outside"] for v in vs))
        self.assertEqual(self.get(f"/api/docs/{it['id']}/versions/{vs[0]['n']}").status_code, 415)   # no text view
        self.assertEqual(self.get(f"/api/docs/{it['id']}/versions/{vs[0]['n']}/file").status_code, 200)
        self.ok(self.post(f"/api/docs/{it['id']}/versions/{vs[0]['n']}/restore"))       # the outside version
        self.assertEqual(self.cells(it["id"]), {"A1": {"v": 1}, "Z9": {"v": "outside"}})

    def test_limits(self):
        it = self.make("Limits")
        etag = self.open(it["id"])["etag"]
        too_many = {"tabs": [grid(f"T{i}") for i in range(11)]}
        r = self.patch(f"/api/docs/{it['id']}", {"etag": etag, "sheet": too_many})
        self.assertEqual(r.status_code, 422)
        self.assertIn("10 tabs", r.text)
        for ref in ("A5001", "CW1"):
            r = self.change(it["id"], [("Sheet1", ref, None, {"v": 1})], etag)
            self.assertEqual(r.status_code, 422, ref)
        for bad in ({"v": "x" * 40000}, {"v": 1, "f": "fancy"}, {"v": 1, "d": 11}, {"v": float("1e309") if False else 1, "al": "up"},
                    {"v": {"e": "#N/A"}}):
            self.assertEqual(self.change(it["id"], [("Sheet1", "A1", None, bad)], etag).status_code, 422, bad)
        for name in ("a/b", "History", "x" * 32, "'quoted'"):
            r = self.patch(f"/api/docs/{it['id']}", {"etag": etag, "sheet": {"tabs": [grid(name)]}})
            self.assertEqual(r.status_code, 422, name)
        r = self.patch(f"/api/docs/{it['id']}", {"etag": etag, "sheet": {"tabs": [grid("A"), grid("a")]}})
        self.assertEqual(r.status_code, 422)
        csv = self.make("C", fmt="csv")
        r = self.patch(f"/api/docs/{csv['id']}", {"etag": self.open(csv["id"])["etag"], "sheet": {"tabs": [grid("A"), grid("B")]}})
        self.assertEqual(r.status_code, 422)
        self.assertIn("Convert to .xlsx", r.text)

    def test_access(self):
        it = self.make("Private", {"tabs": [grid("S", {"A1": {"v": "secret"}})]})
        etag = self.open(it["id"])["etag"]
        self.assertEqual(self.get(f"/api/docs/{it['id']}", MEERA).status_code, 404)
        self.assertEqual(self.change(it["id"], [("S", "A1", None, {"v": 1})], etag, MEERA).status_code, 404)
        self.assertEqual(self.get(f"/api/docs/{it['id']}/export?format=csv", MEERA).status_code, 404)
        self.ok(self.share(it["id"], MEERA, "viewer"))
        doc = self.open(it["id"], MEERA)
        self.assertFalse(doc["canEdit"])
        self.assertEqual(self.change(it["id"], [("S", "A1", None, {"v": 1})], etag, MEERA).status_code, 403)
        self.assertEqual(self.post(f"/api/docs/{it['id']}/convert", {}, MEERA).status_code, 403)
        self.assertEqual(self.get(f"/api/docs/{it['id']}/export?format=xlsx", MEERA).status_code, 200)

    # -- foreign files ----------------------------------------------------------------
    def test_foreign_file_opens_read_only_and_is_never_overwritten(self):
        from openpyxl.comments import Comment
        root = self.root_id()
        data = foreign(lambda wb, ws: setattr(ws["A1"], "comment", Comment("hi", "x")))
        path = os.path.join(os.path.dirname(self.real(self.make("Anchor")["id"])), "From Excel.xlsx")
        write(path, data)
        self.scan()
        found = self.ok(self.get("/api/search?q=From"))["results"]
        node = next(r for r in found if r["name"] == "From Excel.xlsx")
        self.assertEqual(node["kind"], "sheet")
        doc = self.open(node["id"])
        self.assertTrue(doc["readOnly"])
        self.assertFalse(doc["canEdit"])
        self.assertTrue(doc["mayEdit"])
        self.assertIn("Comments or notes", doc["features"])
        r = self.change(node["id"], [("Bills", "C1", None, {"v": 1})], doc["etag"])
        self.assertEqual(r.status_code, 409)
        self.assertTrue(r.json()["detail"]["readOnly"])
        self.assertEqual(read(path, binary=True), data)                          # untouched
        copy = self.ok(self.post(f"/api/docs/{node['id']}/edit-copy", {}), 201)
        self.assertEqual(copy["name"], "From Excel (editable).xlsx")
        cdoc = self.open(copy["id"])
        self.assertFalse(cdoc["readOnly"])
        self.assertEqual(cdoc["sheet"]["tabs"][0]["cells"]["B2"], {"v": 80})
        _ = root

    def test_formula_the_app_cant_calculate_keeps_excels_value(self):
        data = foreign(lambda wb, ws: ws.__setitem__("C1", "=SQRT(16)"))
        data = zip_with(data, {"xl/worksheets/sheet1.xml": lambda raw: raw.replace(b"<f>SQRT(16)</f><v></v>", b"<f>SQRT(16)</f><v>4</v>")
                               .replace(b"<f>SQRT(16)</f><v />", b"<f>SQRT(16)</f><v>4</v>")})
        out = sheet_xlsx.read(data)
        c1 = out["sheet"]["tabs"][0]["cells"]["C1"]
        self.assertEqual(c1, {"v": "=SQRT(16)", "c": 4, "x": 1})
        self.assertIn("Functions the app doesn't have: SQRT", out["features"])

    # -- csv ---------------------------------------------------------------------------
    def test_csv_saves_only_what_changed_and_converts(self):
        it = self.make("List", fmt="csv")
        path = self.real(it["id"])
        write(path, b"Item;Price\r\nMilk;1.20\r\nBread;2\r\n")
        self.scan()
        doc = self.open(it["id"])
        self.assertEqual(doc["sheet"]["tabs"][0]["cells"]["B2"], {"v": "1.20"})
        self.ok(self.change(it["id"], [("Sheet1", "B3", {"v": 2}, {"v": 2.5}), ("Sheet1", "B4", None, {"v": "=SUM(B2:B3)", "c": 2.5})], doc["etag"]))
        self.assertEqual(read(path, binary=True), b"Item;Price\r\nMilk;1.20\r\nBread;2.5\r\n;=SUM(B2:B3)\r\n")
        nid = self.ok(self.post(f"/api/docs/{it['id']}/convert", {}))["id"]
        self.assertEqual(nid, it["id"])
        node = self.ok(self.get(f"/api/nodes/{nid}"))
        self.assertEqual(node["name"], "List.xlsx")
        self.assertFalse(os.path.exists(path))
        doc = self.open(nid)
        self.assertEqual(doc["format"], "xlsx")
        self.assertEqual(doc["sheet"]["tabs"][0]["cells"]["B4"]["v"], "=SUM(B2:B3)")
        vs = self.ok(self.get(f"/api/docs/{nid}/versions"))["versions"]
        r = self.get(f"/api/docs/{nid}/versions/{vs[0]['n']}/file")
        self.assertIn(".csv", r.headers["content-disposition"])                  # the .csv kept under History
        self.assertEqual(self.post(f"/api/docs/{nid}/convert", {}).status_code, 422)
        r = self.post(f"/api/docs/{nid}/versions/{vs[0]['n']}/restore")
        self.assertEqual(r.status_code, 409)                                        # a .csv into an .xlsx: no

    # -- export / import -------------------------------------------------------------------
    def test_export(self):
        it = self.make("Budget", budget())
        r = self.get(f"/api/docs/{it['id']}/export?format=xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r.headers["content-disposition"])
        self.assertEqual(r.headers["x-content-type-options"], "nosniff")
        self.assertEqual(r.content, read(self.real(it["id"]), binary=True))        # the file itself
        r = self.get(f"/api/docs/{it['id']}/export?format=csv&tab=Budget&mode=values")
        self.assertIn('Budget - Budget.csv', r.headers["content-disposition"])
        text = r.content.decode("utf-8-sig")
        lines = text.splitlines()
        self.assertEqual(lines[0], "Item,Cost,,,")
        self.assertEqual(lines[1], "Rent,1200,Rent x,2025-12-09,'=not a formula")   # text guarded, dates as dates
        self.assertIn("#DIV/0!", lines[2])
        r = self.get(f"/api/docs/{it['id']}/export?format=csv&tab=Budget&mode=formulas")
        lines = r.content.decode("utf-8-sig").splitlines()
        self.assertEqual(lines[1], "Rent,1200,\"=CONCAT(A2,\"\" x\"\")\",2025-12-09,'=not a formula")
        self.assertEqual(self.get(f"/api/docs/{it['id']}/export?format=csv&tab=Nope").status_code, 404)
        csv = self.make("C", fmt="csv")
        etag = self.open(csv["id"])["etag"]
        self.ok(self.change(csv["id"], [("Sheet1", "A1", None, {"v": 3})], etag))
        r = self.get(f"/api/docs/{csv['id']}/export?format=xlsx")
        self.assertEqual(sheet_xlsx.read(r.content)["sheet"]["tabs"][0]["cells"]["A1"], {"v": 3})

    def test_import(self):
        r = self.c.post("/api/import?name=Prices.csv", content=b"a,b\n1,2\n", headers=KABIR)
        it = self.ok(r, 201)
        self.assertEqual(it["name"], "Prices.xlsx")
        self.assertEqual(self.cells(it["id"])["B2"], {"v": 2})
        r = self.c.post("/api/import?name=Prices.csv&format=csv", content=b"a;b\n1;2\n", headers=KABIR)
        it = self.ok(r, 201)
        self.assertEqual(it["name"], "Prices.csv")
        self.assertEqual(read(self.real(it["id"]), binary=True), b"a;b\n1;2\n")
        data = foreign(lambda wb, ws: ws.merge_cells("A5:B5"))
        out = self.ok(self.c.post("/api/import?name=Old.xlsx", content=data, headers=KABIR), 201)
        self.assertEqual(out["leftOut"], ["Merged cells"])
        self.assertFalse(self.open(out["id"])["readOnly"])
        self.assertEqual(self.c.post("/api/import?name=x.pdf", content=b"%PDF", headers=KABIR).status_code, 422)
        self.assertEqual(self.c.post("/api/import?name=x.xlsx", content=b"PK nope", headers=KABIR).status_code, 422)
        self.settings({"max_doc_mb": 1})
        self.assertEqual(self.c.post("/api/import?name=x.csv", content=b"a" * (1024 * 1024 + 1), headers=KABIR).status_code, 413)
        folder = self.create("folder", "Shared")
        self.ok(self.share(folder["id"], MEERA, "viewer"))
        r = self.c.post(f"/api/import?name=x.csv&parentId={folder['id']}", content=b"a\n", headers=MEERA)
        self.assertEqual(r.status_code, 403)

    # -- search ---------------------------------------------------------------------------
    def test_search_finds_cells_and_opens_at_them(self):
        it = self.make("Budget", budget())
        res = self.ok(self.get("/api/search?q=Hire"))["results"]
        hit = next(r for r in res if r["id"] == it["id"])
        self.assertEqual(hit["cell"], {"tab": "Car costs", "ref": "A4"})
        self.assertEqual(hit["cellLabel"], "Budget › Car costs › A4")
        self.assertTrue(hit["snippet"].startswith("Budget › Car costs › A4: ⁅Hire⁆"))
        res = self.ok(self.get("/api/search?q=180"))["results"]                     # a formula's shown value
        hit = next(r for r in res if r["id"] == it["id"])
        self.assertEqual(hit["cell"]["ref"], "B4")
        self.assertEqual(self.ok(self.get("/api/search?q=Hire", MEERA))["results"], [])   # access inside search
        self.assertEqual(self.ok(self.get("/api/search?q=Hire&type=sheet"))["total"], 1)
        # a sheet changed outside the app is indexed by the scan (through the worker)
        path = os.path.join(os.path.dirname(self.real(it["id"])), "Outside.xlsx")
        write(path, foreign(lambda wb, ws: ws.__setitem__("D9", "Lawnmower")))
        self.scan()
        res = self.ok(self.get("/api/search?q=lawnmower"))["results"]
        self.assertEqual([(r["name"], r["cell"]) for r in res], [("Outside.xlsx", {"tab": "Bills", "ref": "D9"})])

    def test_settings_new_sheets_and_currency(self):
        me = self.ok(self.get("/api/me"))
        self.assertEqual(me["prefs"]["newSheets"], "xlsx")
        self.assertEqual(me["app"]["currency"], "EUR")
        self.ok(self.put("/api/me/settings", {"newSheets": "csv"}))
        self.assertEqual(self.ok(self.get("/api/me"))["prefs"]["newSheets"], "csv")
        self.assertEqual(self.put("/api/me/settings", {"newSheets": "ods"}).status_code, 422)
        from app import config
        config.set_currency("gbp")
        try:
            self.assertEqual(self.ok(self.get("/api/me"))["app"]["currency"], "GBP")
            it = self.make("Money", {"tabs": [grid("S", {"A1": dict(CUR, v=5)})]})
            from openpyxl import load_workbook
            self.assertEqual(load_workbook(self.real(it["id"]))["S"]["A1"].number_format, '"£"#,##0.00')
        finally:
            config.set_currency("EUR")

    def test_copy_and_trash_keep_sheets_searchable(self):
        it = self.make("Budget", budget())
        copy = self.ok(self.post(f"/api/nodes/{it['id']}/copy", {}), 201)
        res = self.ok(self.get("/api/search?q=Hire"))["results"]
        self.assertEqual({r["id"] for r in res}, {it["id"], copy["id"]})
        self.ok(self.delete(f"/api/nodes/{it['id']}"))
        res = self.ok(self.get("/api/search?q=Hire"))["results"]
        self.assertEqual({r["id"] for r in res}, {copy["id"]})


if __name__ == "__main__":
    unittest.main()
