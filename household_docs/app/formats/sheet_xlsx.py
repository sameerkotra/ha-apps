""".xlsx sheets (SPEC §8.4, §17.10) with openpyxl.

**Reading** is always done by `xlsx_worker.py` in a separate process with a time limit (a file in /share may
come from anywhere): `read(data)` → {"sheet", "readOnly", "features", "owned"}. The worker's findings become the
app's sheet (sheet_model); anything the app couldn't keep when saving — macros, pivot tables, pictures,
comments, data validation, merged cells, links to other workbooks, charts it didn't make, functions it doesn't
have, formatting it has no control for — is listed in `features`, and such a file opens read-only (the app
never overwrites it, §14) with *Save a copy to edit*.

**Writing** (`write(sheet, currency)`) happens in the app (the sheet was checked by sheet_model.clean_sheet):
values, formulas with the values the browser computed (openpyxl writes formulas without them, so they're put
into the sheet XML afterwards — Excel and other apps show the numbers straight away), number formats, bold,
alignment, widths, frozen panes, tabs, conditional colours as Excel conditional formatting, charts as Excel
charts, and the totals row as real formula cells. Custom document properties mark the file as the app's:
`HouseholdDocs=1` and `HouseholdDocsMeta` (JSON: per tab, its charts and totals row) — so on reopening, charts
and totals the app wrote are recognised as its own.
"""
import hashlib
import io
import json
import logging
import os
import re
import subprocess
import sys
import threading
import zipfile
from collections import OrderedDict
from xml.sax.saxutils import escape

from . import sheet_model as M

logger = logging.getLogger("sheets")

WORKER = os.path.join(os.path.dirname(os.path.abspath(__file__)), "xlsx_worker.py")
TIMEOUT = 20                     # seconds of wall-clock time for one file
_slots = threading.BoundedSemaphore(2)
_cache: "OrderedDict[str, dict]" = OrderedDict()
_cache_lock = threading.Lock()
CACHE_SIZE = 12
META_VERSION = 1


class ReadError(Exception):
    """The file can't be opened (message for people)."""


def limits() -> dict:
    return {"cells": M.MAX_CELLS, "rows": M.MAX_ROWS, "cols": M.MAX_COLS, "tabs": M.MAX_TABS,
            "unzipped": 64 * 1024 * 1024, "parts": 4000, "mem": 1024 * 1024 * 1024, "cpu": TIMEOUT + 5}


def run_worker(data: bytes, timeout: float = TIMEOUT) -> dict:
    """The worker's raw findings (or ReadError)."""
    if not data.startswith(b"PK"):
        raise ReadError("This isn't an .xlsx file.")
    with _slots:
        try:
            p = subprocess.run([sys.executable, "-I", WORKER, json.dumps(limits())], input=data,
                               capture_output=True, timeout=timeout, env={"PATH": os.environ.get("PATH", "")})
        except subprocess.TimeoutExpired:
            raise ReadError(f"Reading this .xlsx took longer than {int(timeout)} seconds, so it was stopped.")
    try:
        out = json.loads(p.stdout or b"{}")
    except ValueError:
        out = {}
    if not out:
        logger.warning("The .xlsx reader stopped (exit %s): %s", p.returncode, (p.stderr or b"")[-300:])
        raise ReadError("This .xlsx couldn't be read safely (the reader ran out of memory or time).")
    if not out.get("ok"):
        raise ReadError(out.get("error") or "This .xlsx couldn't be read.")
    return out


def read(data: bytes) -> dict:
    """The sheet in an .xlsx (cached by content, so saving and searching don't parse the same file twice)."""
    key = hashlib.sha256(data).hexdigest()
    with _cache_lock:
        hit = _cache.get(key)
        if hit is not None:
            _cache.move_to_end(key)
            return json.loads(hit)
    out = from_worker(run_worker(data))
    remember(key, out)
    return out


def remember(sha: str, parsed: dict) -> None:
    with _cache_lock:
        _cache[sha] = json.dumps(parsed)
        _cache.move_to_end(sha)
        while len(_cache) > CACHE_SIZE:
            _cache.popitem(last=False)


def clear_cache() -> None:
    with _cache_lock:
        _cache.clear()


# ====================================================================== worker → sheet
def _value(v):
    if isinstance(v, dict):
        if "d" in v:
            return v["d"], True
        if "e" in v:
            return {"e": v["e"] if v["e"] in M.ERRORS else "#VALUE!"}, False
    return v, False


def _cf_rule(raw: dict, rng: str):
    """An Excel conditional format → the app's colour rule, or None when the app has no such rule."""
    fill = _nearest(raw.get("fill"), M.PALETTE)
    text = _nearest(raw.get("font"), M.TEXT_PALETTE)
    if not fill and not text:
        return None
    out = {"range": rng, "fill": fill, "text": text}
    t, op, f = raw.get("type"), raw.get("operator"), raw.get("formula") or []

    def num(x):
        try:
            return float(x) if "." in x or "e" in x.lower() else int(x)
        except (TypeError, ValueError):
            return None
    if t == "cellIs" and op in ("greaterThan", "lessThan") and len(f) == 1 and num(f[0]) is not None:
        out.update(type="gt" if op == "greaterThan" else "lt", a=num(f[0]))
    elif t == "cellIs" and op == "between" and len(f) == 2 and None not in (num(f[0]), num(f[1])):
        out.update(type="between", a=num(f[0]), b=num(f[1]))
    elif t == "containsText" and raw.get("text"):
        out.update(type="contains", a=raw["text"])
    elif t == "top10" and raw.get("rank") and not raw.get("percent"):
        out.update(type="bottom" if raw.get("bottom") else "top", n=int(raw["rank"]))
    elif t == "duplicateValues":
        out.update(type="dup")
    elif t == "expression" and len(f) == 1:
        m = re.fullmatch(r"AND\(ISNUMBER\(\$?[A-Z]+\$?\d+\),\$?[A-Z]+\$?\d+([<>])DATE\((\d+),(\d+),(\d+)\)\)", f[0].replace(" ", ""))
        if not m:
            return None
        serial = M.iso_to_serial(f"{int(m.group(2)):04d}-{int(m.group(3)):02d}-{int(m.group(4)):02d}")
        if serial is None:
            return None
        out.update(type="before" if m.group(1) == "<" else "after", a=serial)
    else:
        return None
    return out


def _nearest(rgb, palette):
    if not rgb or not re.fullmatch(r"[0-9A-Fa-f]{6}", str(rgb)):
        return None
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    best, dist = None, None
    for name, hexv in palette.items():
        x = [int(hexv[i:i + 2], 16) for i in (0, 2, 4)]
        d = (x[0] - r) ** 2 + (x[1] - g) ** 2 + (x[2] - b) ** 2
        if dist is None or d < dist:
            best, dist = name, d
    return best


def _px(width_chars: float) -> int:
    return max(24, min(1000, int(round(width_chars * 7 + 5))))


def from_worker(raw: dict) -> dict:
    features: dict[str, str] = dict(raw.get("features") or {})
    props = raw.get("props") or {}
    owned = props.get("HouseholdDocs") == "1"
    meta = {}
    if owned:
        try:
            meta = json.loads(props.get("HouseholdDocsMeta") or "{}")
        except ValueError:
            meta = {}
    meta_tabs = meta.get("tabs") if isinstance(meta.get("tabs"), dict) else {}
    names = set(raw.get("names") or [])
    if names:
        features["names"] = "Named ranges"
    unknown: set[str] = set()
    tabs = []
    charts_in_file = 0
    charts_known = 0
    for sh in raw.get("sheets", []):
        title = str(sh.get("title") or "Sheet")
        tm = meta_tabs.get(title) if isinstance(meta_tabs.get(title), dict) else {}
        if sh.get("state") not in (None, "visible"):
            features["hidden"] = "Hidden tabs"
        if sh.get("kind") == "chart":
            charts_in_file += sh.get("charts", 0)
            ch = tm.get("chart")
            if sh.get("charts") == 1 and isinstance(ch, dict):
                try:
                    tabs.append({"name": title, "kind": "chart", "chart": M.clean_chart(ch, own_tab=True)})
                    charts_known += 1
                    continue
                except M.SheetError:
                    pass
            features["charts"] = "Charts not made by the app"
            continue
        cells = {}
        excel_cells = 0
        for coord, v, typ, numfmt, bold, align, quote, cached in sh.get("cells", []):
            if M.parse_ref(coord) is None:
                continue
            cell = {}
            val, was_date = _value(v)
            if typ == "f":
                text = M.from_excel_formula(str(v or ""))
                if text.startswith("="):
                    text = text[1:]
                cell["v"] = "=" + text
                cv, _d = _value(cached)
                if cv is not None:
                    cell["c"] = cv
                if not M.can_calculate(text, names):
                    cell["x"] = 1
                    excel_cells += 1
                    unknown |= {f for f in M.functions_in(text) if f not in M.KNOWN_FUNCTIONS}
            elif typ == "e":
                cell["v"] = "=" + (val["e"] if isinstance(val, dict) else "#VALUE!")
            elif val is not None:
                if isinstance(val, float) and val.is_integer() and abs(val) < 1e15:
                    val = int(val)
                cell["v"] = val
                if quote and isinstance(val, str):
                    cell["q"] = 1
                elif isinstance(val, str) and val.startswith("=") and len(val) > 1:
                    cell["q"] = 1
            fm = M.format_from_excel(numfmt)
            if fm is None:
                features["formats"] = "Number formats the app doesn't have (times, scientific …)"
                fm = {}
            if was_date and not fm:
                fm = {"f": "date"}
            cell.update(fm)
            if bold:
                cell["b"] = 1
            if align in M.ALIGNS:
                cell["al"] = align
            elif align not in (None, "general"):
                features["styles"] = "Colours, borders or fonts the app doesn't keep"
            if cell:
                cells[coord] = cell
        if excel_cells:
            features["functions"] = ("Functions the app doesn't have: " + ", ".join(sorted(unknown))) if unknown \
                else "Formulas the app can't calculate (named ranges, table references or arrays)"
        if sh.get("styled"):
            features["styles"] = "Colours, borders or fonts the app doesn't keep"
        if sh.get("merged"):
            features["merged"] = "Merged cells"
        if sh.get("validations"):
            features["validation"] = "Data validation (drop-down lists, checks)"
        if sh.get("images"):
            features["images"] = "Pictures"
        if sh.get("special"):
            features["arrays"] = "Array formulas"
        if sh.get("hiddenRows") or sh.get("hiddenCols"):
            features["hiddenrc"] = "Hidden rows or columns"
        if sh.get("autofilter"):
            features["filter"] = "Filters (AutoFilter)"
        if sh.get("protected"):
            features["protected"] = "Protected tabs"
        cols = {}
        for letter, (width, lo, hi) in (sh.get("cols") or {}).items():
            try:
                first = int(lo) - 1 if lo else M.col_index(letter)
                last = int(hi) - 1 if hi else first
            except (TypeError, ValueError):
                continue
            for ci in range(max(0, first), min(last, M.MAX_COLS - 1) + 1):
                cols[M.col_name(ci)] = _px(width)
        freeze = {"r": 0, "c": 0}
        fz = M.parse_ref(str(sh.get("freeze") or "").replace("$", ""))
        if fz:
            freeze = {"r": min(fz[1], 50), "c": min(fz[0], 20)}
        cond = []
        for rule in sh.get("cf") or []:
            sq = str(rule.get("sqref") or "")
            rng = M.parse_range(sq) if " " not in sq.strip() else None
            r = _cf_rule(rule, M.range_text(*rng)) if rng else None
            if r is None:
                features["cf"] = "Conditional formatting the app doesn't have"
            else:
                try:
                    cond.append(M.clean_rule(r))
                except M.SheetError:
                    features["cf"] = "Conditional formatting the app doesn't have"
        # charts and the totals row the app wrote (recognised through HouseholdDocsMeta)
        n_charts = sh.get("charts", 0)
        charts_in_file += n_charts
        charts = []
        if n_charts:
            listed = tm.get("charts") if isinstance(tm.get("charts"), list) else []
            if len(listed) == n_charts:
                try:
                    charts = [M.clean_chart(c) for c in listed]
                    charts_known += n_charts
                except M.SheetError:
                    charts = []
            if not charts:
                features["charts"] = "Charts not made by the app"
        totals = None
        tt = tm.get("totals")
        if isinstance(tt, dict) and isinstance(tt.get("row"), int) and isinstance(tt.get("funcs"), dict):
            totals = _take_totals(cells, tt)
        tabs.append({"name": title, "kind": "grid", "cells": cells, "cols": cols, "freeze": freeze,
                     "totals": totals, "cond": cond[:50], "charts": charts[:10]})
    if raw.get("chartParts", 0) > charts_in_file or charts_in_file > charts_known:
        features["charts"] = "Charts not made by the app"
    if not any(t["kind"] == "grid" for t in tabs):
        tabs.insert(0, M.empty_sheet()["tabs"][0])
    sheet = {"tabs": tabs}
    try:
        sheet = M.clean_sheet(sheet)
    except M.SheetError as e:
        raise ReadError(f"This .xlsx can't be opened: {e}")
    # clean_sheet drops computed "x" marks only for empty cells; keep the rest as read
    feats = [features[k] for k in sorted(features)]
    return {"sheet": sheet, "readOnly": bool(feats), "features": feats, "owned": owned}


def _take_totals(cells: dict, tt: dict):
    """The totals row the app wrote: the formula cells (and the "Total" label) come out of the grid again."""
    row = tt["row"]
    funcs = {}
    for col, fn in tt["funcs"].items():
        if fn not in M.TOTALS or M.parse_ref(f"{col}{row + 1}") is None:
            return None
        ref = f"{col}{row + 1}"
        cell = cells.get(ref)
        want = f"={fn}({col}1:{col}{row})"
        if not cell or str(cell.get("v", "")).replace("$", "").upper() != want.upper():
            return None
        funcs[col] = fn
    for col in funcs:
        cells.pop(f"{col}{row + 1}", None)
    label = tt.get("label")
    if isinstance(label, str) and M.parse_ref(f"{label}{row + 1}"):
        lc = cells.get(f"{label}{row + 1}")
        if lc and lc.get("v") == "Total":
            cells.pop(f"{label}{row + 1}", None)
    for ref in list(cells):
        p = M.parse_ref(ref)
        if p and p[1] == row and not set(cells[ref]) - {"f", "d", "red", "b", "al"}:
            cells.pop(ref, None)            # formats only, on the totals row
    return funcs or None


# ====================================================================== sheet → .xlsx
def totals_row(tab: dict) -> int:
    """Where the totals row goes: right under the last filled row (0-based)."""
    return M.last_row(tab.get("cells") or {}) + 1


def total_value(fn: str, cells: list):
    """A totals cell's value from the column's numbers (what the formula gives)."""
    xs = []
    for c in cells:
        v = c.get("c") if M.is_formula(c) else c.get("v")
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            xs.append(v)
    if fn == "COUNT":
        return len(xs)
    if not xs:
        return {"e": "#DIV/0!"} if fn == "AVERAGE" else 0
    return {"SUM": sum(xs), "AVERAGE": sum(xs) / len(xs), "MIN": min(xs), "MAX": max(xs)}[fn]


def write(sheet: dict, currency: str | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.packaging.custom import StringProperty
    from openpyxl.styles import Alignment, Font
    from openpyxl.workbook.properties import CalcProperties

    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    aligns = {a: Alignment(horizontal=a) for a in M.ALIGNS}
    meta = {"v": META_VERSION, "tabs": {}}
    cached: dict[int, dict[str, object]] = {}
    sheets = {}
    grid_index = 0
    for tab in sheet["tabs"]:
        if tab["kind"] == "chart":
            sheets[tab["name"]] = wb.create_chartsheet(tab["name"])
            meta["tabs"][tab["name"]] = {"chart": tab["chart"]}
            continue
        grid_index += 1
        ws = wb.create_sheet(tab["name"])
        sheets[tab["name"]] = ws
        values = cached.setdefault(grid_index, {})
        for ref, cell in tab["cells"].items():
            xc = ws[ref]
            v = cell.get("v")
            if M.is_formula(cell):
                xc.value = "=" + M.to_excel_formula(v[1:])
                values[ref] = cell.get("c")
            elif isinstance(v, str) and (cell.get("q") or v.startswith("=")):
                xc.value = v
                xc.data_type = "s"
                xc.quotePrefix = True
            else:
                xc.value = v
            nfmt = M.excel_format(cell, currency)
            if nfmt != "General":
                xc.number_format = nfmt
            if cell.get("b"):
                xc.font = bold
            if cell.get("al"):
                xc.alignment = aligns[cell["al"]]
        for letter, px in (tab.get("cols") or {}).items():
            ws.column_dimensions[letter].width = round(max(1.0, (px - 5) / 7), 2)
        fr = tab.get("freeze") or {}
        if fr.get("r") or fr.get("c"):
            ws.freeze_panes = M.addr(fr.get("c", 0), fr.get("r", 0))
        tm = {}
        if tab.get("totals"):
            row = totals_row(tab)
            funcs = {}
            for col, fn in sorted(tab["totals"].items(), key=lambda x: M.col_index(x[0])):
                ref = f"{col}{row + 1}"
                xc = ws[ref]
                xc.value = f"={fn}({col}1:{col}{row})"
                xc.font = bold
                src = [c for r2, c in tab["cells"].items() if r2.rstrip("0123456789") == col]
                fmt_cell = next((c for c in src if c.get("f") not in (None, "general", "text", "date")), None)
                if fmt_cell and fn != "COUNT":
                    xc.number_format = M.excel_format(fmt_cell, currency)
                values[ref] = total_value(fn, src)
                funcs[col] = fn
            label = None
            for ci in range(M.MAX_COLS):
                letter = M.col_name(ci)
                if letter not in funcs and f"{letter}{row + 1}" not in tab["cells"]:
                    label = letter
                    break
                if letter in funcs:
                    break
            if label is not None and M.col_index(label) < min(M.col_index(c) for c in funcs):
                ws[f"{label}{row + 1}"].value = "Total"
                ws[f"{label}{row + 1}"].font = bold
            tm["totals"] = {"row": row, "funcs": funcs, "label": label}
        if tab.get("cond"):
            _write_cf(ws, tab["cond"])
        if tab.get("charts"):
            tm["charts"] = tab["charts"]
        if tm:
            meta["tabs"][tab["name"]] = tm
    # charts (after every tab exists: a chart tab draws from another tab)
    for tab in sheet["tabs"]:
        if tab["kind"] == "chart":
            src = tab["chart"]["source"]
            ws = next((sheets[n] for n in sheets if n.casefold() == src["tab"].casefold()), None)
            if ws is not None and hasattr(ws, "cell"):
                sheets[tab["name"]].add_chart(_chart(ws, tab["chart"]["type"], src["range"], tab["chart"]["title"]))
        else:
            ws = sheets[tab["name"]]
            below = totals_row(tab) + (2 if tab.get("totals") else 1) + 1
            for i, ch in enumerate(tab.get("charts") or []):
                c = _chart(ws, ch["type"], ch["range"], ch["title"])
                ws.add_chart(c, f"{M.col_name((i % 2) * 9)}{below + (i // 2) * 18 + 1}")
    wb.custom_doc_props.append(StringProperty(name="HouseholdDocs", value="1"))
    wb.custom_doc_props.append(StringProperty(name="HouseholdDocsMeta", value=json.dumps(meta, separators=(",", ":"))))
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    buf = io.BytesIO()
    wb.save(buf)
    return _put_cached_values(buf.getvalue(), cached)


def _write_cf(ws, rules) -> None:
    from openpyxl.formatting.rule import CellIsRule, FormulaRule, Rule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.styles.differential import DifferentialStyle
    for r in rules:
        fill = PatternFill(start_color="FF" + M.PALETTE[r["fill"]], end_color="FF" + M.PALETTE[r["fill"]],
                           fill_type="solid") if r.get("fill") else None
        font = Font(color="FF" + M.TEXT_PALETTE[r["text"]]) if r.get("text") else None
        rng = r["range"]
        top_left = rng.split(":")[0]
        t = r["type"]
        if t in ("gt", "lt"):
            rule = CellIsRule(operator="greaterThan" if t == "gt" else "lessThan", formula=[M.num_text(r["a"])],
                              fill=fill, font=font)
        elif t == "between":
            lo, hi = sorted((r["a"], r["b"]))
            rule = CellIsRule(operator="between", formula=[M.num_text(lo), M.num_text(hi)], fill=fill, font=font)
        elif t in ("before", "after"):
            y, mo, d = M.serial_to_iso(r["a"]).split("-")
            rule = FormulaRule(formula=[f"AND(ISNUMBER({top_left}),{top_left}{'<' if t == 'before' else '>'}"
                                        f"DATE({int(y)},{int(mo)},{int(d)}))"], fill=fill, font=font)
        elif t == "contains":
            text = r["a"].replace('"', '""')
            rule = Rule(type="containsText", operator="containsText", text=r["a"],
                        dxf=DifferentialStyle(fill=fill, font=font),
                        formula=[f'NOT(ISERROR(SEARCH("{text}",{top_left})))'])
        elif t in ("top", "bottom"):
            rule = Rule(type="top10", rank=r["n"], bottom=t == "bottom", dxf=DifferentialStyle(fill=fill, font=font))
        else:
            rule = Rule(type="duplicateValues", dxf=DifferentialStyle(fill=fill, font=font))
        ws.conditional_formatting.add(rng, rule)


SERIES_COLOURS = ("1A73E8", "E8710A", "1E8E3E", "D93025", "9334E6", "12B5CB", "F9AB00", "5F6368")


def _chart(ws, kind: str, rng: str, title: str):
    """An Excel chart of a range: the first row names the series, the first column holds the categories."""
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.chart.series import SeriesLabel  # noqa: F401
    c1, r1, c2, r2 = M.parse_range(rng)
    if kind == "pie":
        ch = PieChart()
    elif kind == "line":
        ch = LineChart()
    else:
        ch = BarChart()
        ch.type = "col"
    if title:
        ch.title = title
    if c2 > c1:
        data = Reference(ws, min_col=c1 + 2, min_row=r1 + 1, max_col=c2 + 1 if kind != "pie" else c1 + 2, max_row=r2 + 1)
        cats = Reference(ws, min_col=c1 + 1, min_row=r1 + 2, max_row=r2 + 1)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
    else:
        data = Reference(ws, min_col=c1 + 1, min_row=r1 + 1, max_row=r2 + 1)
        ch.add_data(data, titles_from_data=True)
    if kind != "pie":
        for i, s in enumerate(ch.series):
            colour = SERIES_COLOURS[i % len(SERIES_COLOURS)]
            if kind == "line":
                s.graphicalProperties.line.solidFill = colour
            else:
                s.graphicalProperties.solidFill = colour
                s.graphicalProperties.line.solidFill = colour
    ch.width, ch.height = 16, 8
    return ch


_FORMULA_CELL = re.compile(r'<c r="([A-Z]{1,3}[0-9]{1,7})"((?: [A-Za-z:]+="[^"]*")*)><f>(.*?)</f><v\s*/></c>', re.S)


def _put_cached_values(data: bytes, cached: dict[int, dict]) -> bytes:
    """openpyxl writes <f>…</f><v/>: put the computed value into <v> (and the type into t=)."""
    if not any(cached.values()):
        return data
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            raw = zin.read(info.filename)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", info.filename)
            if m and cached.get(int(m.group(1))):
                values = cached[int(m.group(1))]

                def fill(mt):
                    ref, attrs, f = mt.group(1), mt.group(2), mt.group(3)
                    v = values.get(ref)
                    attrs = re.sub(r' t="[^"]*"', "", attrs)
                    if v is None or (isinstance(v, dict) and v.get("e") == "#CYCLE!"):
                        return mt.group(0)
                    if isinstance(v, bool):
                        return f'<c r="{ref}"{attrs} t="b"><f>{f}</f><v>{1 if v else 0}</v></c>'
                    if isinstance(v, (int, float)):
                        return f'<c r="{ref}"{attrs}><f>{f}</f><v>{M.num_text(v) if not isinstance(v, int) else v}</v></c>'
                    if isinstance(v, dict):
                        return f'<c r="{ref}"{attrs} t="e"><f>{f}</f><v>{escape(v.get("e", "#VALUE!"))}</v></c>'
                    return f'<c r="{ref}"{attrs} t="str"><f>{f}</f><v>{escape(str(v))}</v></c>'
                raw = _FORMULA_CELL.sub(fill, raw.decode("utf-8")).encode("utf-8")
            zout.writestr(info, raw)
    return out.getvalue()
